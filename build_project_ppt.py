from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt


ROOT = Path(__file__).parent
TEMPLATE = ROOT / "AICTE_IBM_BOB_Project_Submission_Template_for_EduentFoundation_.pptx"
OUTPUT = ROOT / "AICTE_IBM_Mock_Interview_Trainer_Project.pptx"


def remove_pictures(slide):
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shape._element.getparent().remove(shape._element)


def set_text(shape, text, font_size=18, bold=False):
    if not hasattr(shape, "text_frame"):
        return
    shape.text_frame.clear()
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = bold


def replace_named(slide, name, text, font_size=18, bold=False):
    for shape in slide.shapes:
        if shape.name == name:
            set_text(shape, text, font_size, bold)
            return True
    return False


def add_content(slide, text, left=0.8, top=1.4, width=11.7, height=5.5, font_size=18):
    box = slide.shapes.add_textbox(
        Pt(left * 72), Pt(top * 72), Pt(width * 72), Pt(height * 72)
    )
    set_text(box, text, font_size)
    return box


def main():
    presentation = Presentation(str(TEMPLATE))

    set_text(presentation.slides[0].shapes[0], "IBM University Engagement Project Submission", 28, True)
    set_text(presentation.slides[0].shapes[1], "AI Mock Interview Trainer Agent", 24, True)

    slide = presentation.slides[1]
    replace_named(slide, "TextBox 3", "Domain of Project - Education / Artificial Intelligence", 20, True)
    replace_named(slide, "TextBox 6", "Project Title - AI Mock Interview Trainer Agent", 20, True)
    replace_named(
        slide,
        "TextBox 2",
        "Student Name: Vaishnavi Nikam\nEmail ID: [add email]\nMobile / WhatsApp: [add number]\nTarget Role: AI Engineer / Machine Learning Engineer",
        18,
    )

    content = {
        2: ("Problem Statement: AI Mock Interview Trainer", "The Challenge\n\nCandidates often lack access to realistic, personalized interview practice. Generic question lists do not understand a candidate's resume, projects, target role, or answer quality. This makes it difficult to identify knowledge gaps, improve communication, and build confidence before a real interview.\n\nThe Objective\n\nBuild an AI-powered interview trainer that analyzes a resume, generates role-specific questions, supports text and voice answers, evaluates responses, and maintains interview history.\n\nKey needs\n- Resume-aware and personalized questions\n- Technical, project, HR, and coding practice\n- Retrieval-Augmented Generation using an interview knowledge base\n- Instant feedback on submitted answers\n- A simple interface for repeated practice"),
        3: ("Proposed Solution: AI Mock Interview Trainer", "The proposed solution is a Streamlit-based AI interview practice platform powered by IBM watsonx.ai and Granite models.\n\nThe workflow begins with resume upload and parsing. The candidate profile is then used with the selected role, difficulty, question count, and answer style to generate a structured mock interview.\n\nA RAG pipeline retrieves relevant interview concepts from the local knowledge base. Granite generates questions, evaluates answers, and powers resume-aware chat practice. Candidates can answer by typing or recording voice, and completed sessions can be saved for later review.\n\nCore modules\n- Resume Intelligence: extracts skills, projects, and profile text\n- Interview Generator: creates structured role-specific questions\n- Answer Evaluator: returns actionable feedback\n- Chat Trainer: provides interactive preparation support\n- Voice Module: converts spoken answers to text\n- History Module: stores completed interviews"),
        4: ("Technology Used", "IBM watsonx.ai / Granite Models\n- Question generation, answer evaluation, and conversational coaching\n\nIBM Cloud\n- Secure access to hosted foundation models through IBM watsonx.ai\n\nPython\n- Core application and AI orchestration language\n\nStreamlit\n- Interactive web interface for resume analysis and mock interviews\n\nRAG\n- Retrieves relevant interview knowledge before generation\n\nChromaDB\n- Vector storage for knowledge-base retrieval\n\nLangChain and Hugging Face\n- Retrieval and embedding ecosystem\n\nPyPDF2 and python-docx\n- PDF and DOCX resume extraction\n\nSpeechRecognition\n- Voice answer transcription"),
        5: ("Required Files Created for the Project", "app_ui.py - Streamlit application interface\nagent.py - IBM Granite generation, evaluation, and chat functions\nresume_parser.py - Resume extraction and profile analysis\nrag.py / rag_agent.py - Retrieval-Augmented Generation logic\nvector_store.py - Vector database operations\nhistory.py - Interview history persistence\nvoice.py - Speech-to-text support\nknowledge_base/ - Interview preparation documents\nrequirements.txt - Python dependencies\nREADME.md - Setup and usage instructions\n.env - Local credentials only; excluded from Git"),
        6: ("Application Screens and User Flow", "1. Home: access the interview trainer modules\n2. Resume: upload a PDF or DOCX and review the extracted profile\n3. Chat Trainer: ask resume-aware preparation questions\n4. Mock Interview: choose role, difficulty, style, and question count\n5. Answer: respond using text or voice transcription\n6. Feedback: receive AI evaluation after each response\n7. History: save completed interview sessions for later review\n\nThe interface is implemented in app_ui.py with Streamlit session state coordinating the active profile, questions, answers, and interview progress."),
        7: ("Architecture Blueprint", "Candidate\n  |\n  v\nStreamlit UI (app_ui.py)\n  |\n  +--> Resume Parser --> Candidate Profile\n  |\n  +--> Interview Agent --> Granite Model\n  |                         ^\n  +--> Answer Evaluator ---+\n  |\n  +--> RAG Agent --> ChromaDB Vector Store --> Knowledge Base\n  |\n  +--> Voice Module --> Transcribed Answer\n  |\n  +--> History Module --> Local SQLite Database\n\nIBM watsonx.ai provides the hosted Granite model inference layer. The local modules manage parsing, retrieval, interaction, and persistence."),
        8: ("Role of Agentic AI in the Solution", "Agentic AI enables the trainer to coordinate several focused capabilities instead of behaving like a single static question generator.\n\n- Resume agent: understands the candidate's background\n- Interview agent: creates the next relevant question\n- Retrieval agent: finds supporting interview knowledge\n- Evaluation agent: analyzes the submitted answer\n- Chat agent: responds to preparation requests\n- Voice capability: converts spoken answers into usable text\n\nTogether, these components create a context-aware and goal-driven practice loop. The system adapts questions and feedback to the candidate profile, target role, selected difficulty, and current answer."),
        9: ("Token Usage and Prompt Design", "Prompt inputs include:\n- Target role and interview difficulty\n- Candidate skills, projects, and resume text\n- Question type and answer style\n- Retrieved knowledge-base context\n- Candidate's submitted answer\n\nOutput includes:\n- Structured interview questions\n- Evaluation feedback and improvement guidance\n- Resume-aware coaching responses\n\nToken usage depends on resume length, retrieved context, number of questions, and answer size. Keeping prompts focused and requesting structured JSON helps control cost and makes the Streamlit interface reliable."),
        10: ("Project Output: Resume Intelligence", "Expected application output\n\nResume upload\n  -> PDF / DOCX text extraction\n  -> Skills and projects identified\n  -> Candidate profile displayed\n  -> Profile can be edited and saved\n\nThe Resume page gives the interview generator and chat trainer a consistent candidate context. This allows the system to ask questions about the candidate's actual projects instead of relying only on generic prompts."),
        11: ("Project Output: Mock Interview Setup", "The Mock Interview page allows the candidate to configure:\n\n- Target role: AI Engineer, ML Engineer, Data Analyst, Software Developer, or Data Scientist\n- Question count: 1 to 10\n- Difficulty: Easy, Medium, or Hard\n- Answer style: Short, Medium, or Detailed\n- Answer method: Text or Voice\n\nAfter starting, the interface displays one question at a time with progress tracking and Submit / Skip controls."),
        12: ("Project Output: Answer Evaluation", "For each submitted answer, the system sends the question and answer to the IBM Granite evaluator.\n\nFeedback can identify:\n- What the candidate explained well\n- Missing technical concepts\n- Clarity and relevance issues\n- Specific ways to improve the response\n\nThe result is stored with the question and answer so the candidate can review the full interview summary after completion."),
        13: ("Project Output: Chat Trainer", "The Chat Trainer provides resume-aware preparation through natural-language prompts.\n\nExample prompts:\n- Generate interview questions from my projects\n- Review my resume for an AI Engineer role\n- Ask difficult ML interview questions\n- Improve my interview answers\n- Prepare HR questions\n\nChat history remains available during the session, and the candidate can clear it when starting a new practice topic."),
        14: ("Novelty and Uniqueness", "The project combines personalized resume intelligence, RAG, interactive practice, voice input, and answer evaluation in one workflow.\n\nResume-grounded personalization - questions are connected to the candidate's skills and projects.\n\nMulti-modal practice - candidates can type or speak their answers.\n\nContinuous feedback loop - every response produces guidance for improvement.\n\nRetrieval-supported generation - the knowledge base supplies relevant interview concepts.\n\nSession continuity - completed interviews can be saved and reviewed.\n\nThe result is a practical interview companion rather than a static collection of questions."),
        15: ("GitHub Repository", "Public repository:\nhttps://github.com/30nikamvaishnavi/AI-Mock-Interview-Trainer-Agent\n\nThe repository contains:\n- Application source files\n- Knowledge-base documents\n- requirements.txt\n- README.md with setup instructions\n- Test files\n\nSensitive credentials in .env, local virtual environments, generated databases, uploaded resumes, audio files, and reports are excluded from the repository."),
        16: ("Future Scope", "1. Cloud deployment\nDeploy the Streamlit application with managed secrets and a hosted vector database.\n\n2. Analytics dashboard\nTrack scores, topic-level weaknesses, progress over time, and repeated-question performance.\n\n3. Better voice experience\nAdd browser-native recording, confidence signals, pronunciation analysis, and multilingual support.\n\n4. More input formats\nSupport webcam interview practice, coding editor input, and richer resume formats.\n\n5. Adaptive interviews\nUse previous feedback to change the next question difficulty and focus areas automatically."),
        17: ("Certificates and Learning", "Relevant learning areas demonstrated by the project:\n\n- IBM watsonx.ai and Granite model integration\n- Retrieval-Augmented Generation\n- LangChain and vector databases\n- Streamlit application development\n- Resume parsing and document processing\n- Speech-to-text interaction\n\nAdd verified Credly, IBM, or program certificates here before final submission."),
        18: ("Project Demonstration", "The completed project can be demonstrated through the following workflow:\n\n1. Launch the Streamlit app\n2. Upload and analyze a resume\n3. Open Chat Trainer and ask a preparation question\n4. Start a mock interview\n5. Submit a text or voice answer\n6. Review AI feedback\n7. Save the completed interview\n\nAdd screenshots from the running application here for the final submission version."),
    }

    for slide_number, (title, body) in content.items():
        slide = presentation.slides[slide_number]
        remove_pictures(slide)
        text = f"{title}\n\n{body}"
        text_shapes = [s for s in slide.shapes if hasattr(s, "text_frame")]
        if text_shapes:
            set_text(text_shapes[0], text, 16)
            for extra in text_shapes[1:]:
                if extra.name not in {"Google Shape;278;p18", "Google Shape;279;p18"}:
                    set_text(extra, "", 16)
        else:
            add_content(slide, text)

    # The last template slide is intentionally retained as a clean closing slide.
    remove_pictures(presentation.slides[19])
    add_content(presentation.slides[19], "Thank You\n\nAI Mock Interview Trainer Agent", 2.0, 2.4, 9.5, 2.0, 28)

    presentation.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()